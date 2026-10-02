from __future__ import annotations

import csv
from collections.abc import Iterable
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.services.ingestion_contracts import (
    AcquiredSource,
    AdapterCapabilities,
    ContentUnit,
    ExtractionResult,
    ExtractionWarning,
    IngestionContractError,
    IngestionExecutionControl,
    IngestionUnitType,
    ResolvedAdapterConfiguration,
    ValidationIssue,
    ValidationResult,
)


class SpreadsheetIngestionAdapter:
    """Generic XLSX and CSV structure extractor.

    The adapter discovers sheets and rows without fixed sheet names, business
    fields or customer-specific profiles. It returns structured row units only.
    It does not persist data, publish artifacts, execute formulas externally or
    invoke AI components.
    """

    adapter_key = "platform.spreadsheet.generic"
    adapter_version = "1.0.0"
    capabilities = AdapterCapabilities(
        supported_media_types=frozenset(
            {
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "text/csv",
            }
        ),
        priority=100,
        requires_external_binary=False,
        supports_ocr=False,
        supports_structured_output=True,
        supports_cancellation=True,
        community_available=True,
        enterprise_available=True,
    )

    def validate(
        self,
        source: AcquiredSource,
        configuration: ResolvedAdapterConfiguration,
    ) -> ValidationResult:
        issues: list[ValidationIssue] = []
        path = Path(source.local_path)

        if source.detected_media_type not in self.capabilities.supported_media_types:
            issues.append(
                ValidationIssue(
                    code="unsupported_media_type",
                    message="detected media type is not supported by the spreadsheet adapter",
                )
            )

        if not path.is_file():
            issues.append(
                ValidationIssue(
                    code="source_not_found",
                    message="acquired source file is not available",
                )
            )
            return ValidationResult(False, source.detected_media_type, tuple(issues))

        if path.stat().st_size != source.content_length:
            issues.append(
                ValidationIssue(
                    code="content_length_mismatch",
                    message="acquired source length does not match the registered source",
                )
            )

        if _sha256_file(path) != source.checksum_sha256:
            issues.append(
                ValidationIssue(
                    code="checksum_mismatch",
                    message="acquired source checksum does not match the registered source",
                )
            )

        if path.stat().st_size == 0:
            issues.append(ValidationIssue(code="empty_content", message="spreadsheet source is empty"))

        max_sheets = _positive_int(configuration.settings, "max_sheets", 200)
        try:
            if source.detected_media_type == "text/csv":
                _read_csv_rows(path, configuration.settings, limit=1)
            else:
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    if not workbook.sheetnames:
                        issues.append(ValidationIssue(code="empty_workbook", message="workbook contains no sheets"))
                    elif len(workbook.sheetnames) > max_sheets:
                        issues.append(
                            ValidationIssue(
                                code="sheet_limit_exceeded",
                                message="workbook exceeds configured sheet limit",
                            )
                        )
                finally:
                    workbook.close()
        except (InvalidFileException, OSError, UnicodeDecodeError, csv.Error, ValueError):
            issues.append(
                ValidationIssue(
                    code="invalid_spreadsheet",
                    message="source cannot be parsed as a supported spreadsheet",
                )
            )

        return ValidationResult(not issues, source.detected_media_type, tuple(issues))

    def extract(
        self,
        source: AcquiredSource,
        configuration: ResolvedAdapterConfiguration,
        control: IngestionExecutionControl,
    ) -> ExtractionResult:
        validation = self.validate(source, configuration)
        if not validation.accepted:
            raise IngestionContractError("source failed spreadsheet adapter validation")

        control.raise_if_cancellation_requested()
        control.pulse()

        max_rows = _positive_int(configuration.settings, "max_rows", 100_000)
        max_columns = _positive_int(configuration.settings, "max_columns", 1_000)
        header_row_number = _positive_int(configuration.settings, "header_row", 1)
        include_empty_rows = _bool_setting(configuration.settings, "include_empty_rows", False)

        if source.detected_media_type == "text/csv":
            sheets = (("CSV", _read_csv_rows(Path(source.local_path), configuration.settings)),)
            detected_format = "csv"
        else:
            sheets = _read_xlsx_sheets(Path(source.local_path))
            detected_format = "xlsx"

        units: list[ContentUnit] = []
        warnings: list[ExtractionWarning] = []
        total_rows = 0
        total_sheets = 0

        for sheet_name, rows in sheets:
            total_sheets += 1
            control.raise_if_cancellation_requested()

            row_list = list(rows)
            if not row_list:
                warnings.append(
                    ExtractionWarning(
                        code="empty_sheet",
                        message=f"sheet {total_sheets} contains no rows",
                    )
                )
                continue

            if header_row_number > len(row_list):
                raise IngestionContractError("configured header row exceeds available rows")

            headers = _normalize_headers(row_list[header_row_number - 1], max_columns)
            if not headers:
                raise IngestionContractError("spreadsheet header row contains no usable columns")

            for physical_row_number, values in enumerate(row_list[header_row_number:], start=header_row_number + 1):
                control.raise_if_cancellation_requested()
                if total_rows >= max_rows:
                    raise IngestionContractError("spreadsheet exceeds configured row limit")

                normalized_values = tuple(_normalize_cell(value) for value in values[:max_columns])
                if not include_empty_rows and all(value is None for value in normalized_values):
                    continue

                record = {
                    header: normalized_values[index] if index < len(normalized_values) else None
                    for index, header in enumerate(headers)
                }
                canonical = _canonical_record(record)
                content_hash = sha256(canonical.encode("utf-8")).hexdigest()

                units.append(
                    ContentUnit(
                        unit_key=(f"sheet:{_stable_name(sheet_name)}:row:{physical_row_number}:{content_hash[:16]}"),
                        ordinal=len(units),
                        unit_type=IngestionUnitType.ROW,
                        content_hash=content_hash,
                        structured_data=record,
                        source_locator={
                            "sheet_name": sheet_name,
                            "row_number": physical_row_number,
                        },
                        attributes={
                            "column_count": len(headers),
                            "header_row": header_row_number,
                        },
                    )
                )
                total_rows += 1
                control.pulse()

        if not units:
            raise IngestionContractError("spreadsheet contains no extractable data rows")

        return ExtractionResult(
            adapter_key=self.adapter_key,
            adapter_version=self.adapter_version,
            detected_media_type=source.detected_media_type,
            detected_format=detected_format,
            content_units=tuple(units),
            proposed_artifacts=(),
            metrics={
                "source_bytes": source.content_length,
                "sheet_count": total_sheets,
                "data_rows": total_rows,
                "content_units": len(units),
                "formulas_evaluated": False,
            },
            warnings=tuple(warnings),
            quality_observations=(),
        )


def _read_xlsx_sheets(path: Path) -> tuple[tuple[str, Iterable[tuple[Any, ...]]], ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return tuple(
            (worksheet.title, tuple(worksheet.iter_rows(values_only=True))) for worksheet in workbook.worksheets
        )
    finally:
        workbook.close()


def _read_csv_rows(
    path: Path,
    settings: Any,
    *,
    limit: int | None = None,
) -> tuple[tuple[str, ...], ...]:
    encoding = str(settings.get("encoding", "utf-8-sig")).strip().lower()
    delimiter = str(settings.get("delimiter", ","))
    if len(delimiter) != 1:
        raise IngestionContractError("delimiter must contain exactly one character")

    rows: list[tuple[str, ...]] = []
    with path.open("r", encoding=encoding, newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row in reader:
            rows.append(tuple(row))
            if limit is not None and len(rows) >= limit:
                break
    return tuple(rows)


def _normalize_headers(values: tuple[Any, ...], max_columns: int) -> tuple[str, ...]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values[:max_columns], start=1):
        base = str(value).strip() if value is not None else ""
        if not base:
            base = f"column_{index}"
        count = counts.get(base, 0) + 1
        counts[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return tuple(headers)


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        normalized = value.replace("\r\n", "\n").replace("\r", "\n").strip()
        return normalized or None
    if isinstance(value, int | float | bool):
        return value
    return str(value)


def _canonical_record(record: dict[str, Any]) -> str:
    return "\n".join(f"{key}={record[key]!r}" for key in sorted(record))


def _stable_name(value: str) -> str:
    normalized = "-".join(value.strip().lower().split())
    return normalized or "sheet"


def _sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _positive_int(settings: Any, key: str, default: int) -> int:
    value = settings.get(key, default)
    if not isinstance(value, int) or isinstance(value, bool) or value <= 0:
        raise IngestionContractError(f"{key} must be a positive integer")
    return value


def _bool_setting(settings: Any, key: str, default: bool) -> bool:
    value = settings.get(key, default)
    if not isinstance(value, bool):
        raise IngestionContractError(f"{key} must be a boolean")
    return value
